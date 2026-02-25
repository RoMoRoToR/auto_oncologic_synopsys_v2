from __future__ import annotations

import io
from typing import Any, Dict, List, Optional

from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.shared import OxmlElement, qn


def _as_list(x: Any) -> List[str]:
    if x is None:
        return []
    if isinstance(x, list):
        return [str(v).strip() for v in x if str(v).strip()]
    s = str(x).strip()
    return [s] if s else []


def _bullets(items: List[str]) -> str:
    if not items:
        return ""
    return "\n".join([f"• {i}" for i in items])


def _set_cell_shading(cell, fill: str = "F5F7FA"):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def build_synopsis_docx(
    synopsis: Dict[str, Any],
    decision: Optional[Dict[str, Any]] = None,
    stats: Optional[Dict[str, Any]] = None,
    timeline: Optional[Dict[str, Any]] = None,
    rag_summary: Optional[Dict[str, Any]] = None,
) -> bytes:
    decision = decision or {}
    stats = stats or {}
    timeline = timeline or {}
    rag_summary = rag_summary or {}

    doc = Document()

    # Base typography
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(11)

    # Title
    title = doc.add_paragraph("СИНОПСИС ПРОТОКОЛА")
    title.runs[0].bold = True
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph("")

    table = doc.add_table(rows=0, cols=2)
    table.autofit = False
    table.columns[0].width = Cm(6.0)
    table.columns[1].width = Cm(11.5)

    def add_row(label: str, value: str, shade_left: bool = True):
        row = table.add_row()
        left = row.cells[0]
        right = row.cells[1]
        left.text = label
        if shade_left:
            _set_cell_shading(left, "F3F5F7")
        for p in left.paragraphs:
            for r in p.runs:
                r.bold = True
        right.text = value or ""

    inn = str(synopsis.get("activeSubstance") or "")
    dose_form = str(synopsis.get("dosageForm") or "")
    prot_title = str(synopsis.get("protocolTitle") or "")
    prot_id = str(synopsis.get("protocolNumber") or "")

    add_row("Название протокола:", prot_title)
    add_row("Идентификационный номер протокола:", prot_id)
    add_row("Спонсор исследования:", str(synopsis.get("sponsor") or ""))
    add_row("Исследовательский центр:", str(synopsis.get("clinicalCenter") or ""))
    add_row("Биоаналитическая лаборатория:", str(synopsis.get("bioanalyticalLab") or ""))
    add_row("Фаза клинического исследования:", str(synopsis.get("clinicalPhase") or ""))
    add_row("Название исследуемого препарата:", str(synopsis.get("investigationalProduct") or ""))
    add_row("Действующее вещество:", inn)
    add_row("Лекарственная форма и дозировка:", dose_form)

    add_row("Цель исследования:", str(synopsis.get("objectives") or ""))
    add_row("Задачи исследования:", _bullets(_as_list(synopsis.get("tasks"))))

    design_block = str(synopsis.get("design") or "")
    periods = str(synopsis.get("studyPeriods") or "")
    duration = str(synopsis.get("duration") or "")
    washout = decision.get("washout_days")
    if washout:
        periods = (periods + f"; washout {washout} дней").strip("; ")
    add_row("Дизайн исследования:", design_block)
    add_row("Периоды / последовательности:", periods)
    add_row("Длительность:", duration)

    add_row("Популяция:", str(synopsis.get("population") or ""))
    add_row("Критерии включения:", _bullets(_as_list(synopsis.get("inclusionCriteria"))))
    add_row("Критерии исключения:", _bullets(_as_list(synopsis.get("exclusionCriteria"))))
    add_row("Критерии прекращения участия:", _bullets(_as_list(synopsis.get("withdrawalCriteria"))))

    add_row("Режим дозирования (Test):", str(synopsis.get("testProductRegimen") or ""))
    add_row("Режим дозирования (Reference):", str(synopsis.get("referenceProductRegimen") or ""))

    tps = timeline.get("timepoints_h") if isinstance(timeline.get("timepoints_h"), list) else []
    tp_str = ", ".join(str(x) for x in tps[:40]) + ("…" if len(tps) > 40 else "")
    meth = str(synopsis.get("methodology") or "")
    if tp_str:
        meth = (meth + f"\n\nТочки отбора (ч): {tp_str}").strip()
    add_row("Отбор проб / методология:", meth)
    add_row("Биоаналитика:", str(synopsis.get("analyticalMethod") or ""))

    add_row("PK параметры:", str(synopsis.get("pkParameters") or ""))
    add_row("Критерии биоэквивалентности:", str(synopsis.get("beCriteria") or ""))
    add_row("Расчёт выборки:", str(synopsis.get("sampleSizeCalculation") or ""))

    add_row("Безопасность:", str(synopsis.get("safetyAnalysis") or ""))
    add_row("Этические аспекты:", str(synopsis.get("ethicalAspects") or ""))
    add_row("Дата версии:", str(synopsis.get("versionDate") or ""))

    bib = synopsis.get("bibliography") if isinstance(synopsis.get("bibliography"), list) else []
    if bib:
        items = []
        for b in bib[:20]:
            if isinstance(b, dict):
                t = (b.get("title") or "").strip()
                u = (b.get("uri") or "").strip()
                if t and u:
                    items.append(f"{t} — {u}")
                elif t:
                    items.append(t)
            else:
                items.append(str(b))
        add_row("Источники:", _bullets(items))
    else:
        add_row("Источники:", "")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
