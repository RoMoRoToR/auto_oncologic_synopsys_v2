from __future__ import annotations

import os
import re
from dataclasses import dataclass
from datetime import datetime, date
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import openpyxl


def _norm(s: Any) -> str:
    t = "" if s is None else str(s)
    t = t.strip().lower().replace("ё", "е")
    t = re.sub(r"\s+", " ", t)
    return t


def _parse_dose_mg(dosage: str) -> Optional[float]:
    s = _norm(dosage)
    if not s:
        return None
    s = s.replace(",", ".")
    m = re.search(r"([0-9]+(?:\.[0-9]+)?)", s)
    if not m:
        return None
    try:
        return float(m.group(1))
    except ValueError:
        return None


def _date_to_iso(x: Any) -> Optional[str]:
    if x is None:
        return None
    if isinstance(x, (datetime, date)):
        return x.isoformat()
    s = str(x).strip()
    if not s:
        return None
    for fmt in ("%d.%m.%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(s, fmt).date().isoformat()
        except Exception:
            pass
    return None


def _date_sort_key(iso: Optional[str]) -> Tuple[int, str]:
    if not iso:
        return (1, "9999-12-31")
    return (0, iso)


def _form_match(form_text: str, wanted_form: str) -> bool:
    ft = _norm(form_text)
    wf = _norm(wanted_form)
    if not wf:
        return True
    syn = {
        "таблетки": ["таблет", "табл"],
        "капсулы": ["капсул", "caps"],
        "раствор": ["раств", "solution"],
        "суспензия": ["сусп", "susp"],
        "порошок": ["порош", "powder"],
    }
    for k, variants in syn.items():
        if k in wf:
            return any(v in ft for v in variants)
    first = wf.split(" ")[0]
    return first in ft


def _dose_match(form_text: str, dose_mg: Optional[float]) -> bool:
    if dose_mg is None:
        return True
    ft = _norm(form_text)
    if abs(dose_mg - int(dose_mg)) < 1e-9:
        d = str(int(dose_mg))
    else:
        d = str(dose_mg).replace(".", "\\.")
    pat = rf"(^|[^0-9]){d}\s*(мг|mg)\b"
    return re.search(pat, ft) is not None


@dataclass
class GrlsRecord:
    reg_no: str
    reg_date_iso: Optional[str]
    trade_name: str
    inn: str
    forms: str


class ReferenceResolver:
    """
    GRLS Excel:
      I = торговое наименование
      J = МНН
      K = формы выпуска
      D = дата регистрации
      C = номер РУ
    """

    def __init__(self, xlsx_path: str | None = None, sheet_name: str = "Действующий"):
        self.xlsx_path = Path(xlsx_path or os.getenv("GRLS_XLSX_PATH", "data/grls.xlsx"))
        self.sheet_name = sheet_name

        # Do not fail hard if file is missing; return empty list instead.
        self._records: List[GrlsRecord] = []
        self._loaded = False
        self._last_mtime: float | None = None

    def _score(self, rec: GrlsRecord, inn_q: str, dose_mg: Optional[float], form: str) -> int:
        score = 0
        rec_inn = _norm(rec.inn)
        rec_forms = _norm(rec.forms)
        if inn_q:
            if rec_inn == inn_q:
                score += 12
            elif inn_q in rec_inn:
                score += 8
        if _dose_match(rec_forms, dose_mg):
            score += 8
        if _form_match(rec_forms, form):
            score += 6
        return score

    def _load(self) -> None:
        if self._loaded:
            return
        if not self.xlsx_path.exists():
            self._records = []
            self._loaded = True
            return

        mtime = self.xlsx_path.stat().st_mtime
        if self._last_mtime == mtime and self._loaded:
            return

        wb = openpyxl.load_workbook(self.xlsx_path, read_only=True, data_only=True)
        if self.sheet_name not in wb.sheetnames:
            raise ValueError(f"Sheet not found: {self.sheet_name}. Available: {wb.sheetnames}")
        ws = wb[self.sheet_name]

        for r in range(6, ws.max_row + 1):
            reg_no = ws.cell(r, 3).value
            reg_date = ws.cell(r, 4).value
            trade = ws.cell(r, 9).value
            inn = ws.cell(r, 10).value
            forms = ws.cell(r, 11).value

            if not trade or not inn or not forms:
                continue

            rec = GrlsRecord(
                reg_no=str(reg_no or "").strip(),
                reg_date_iso=_date_to_iso(reg_date),
                trade_name=str(trade).strip(),
                inn=str(inn).strip(),
                forms=str(forms).strip(),
            )
            self._records.append(rec)

        self._last_mtime = mtime
        self._loaded = True

    def find_reference_options(
        self,
        inn: str,
        dosage: str,
        dosage_form: str,
        limit: int = 10,
    ) -> Dict[str, Any]:
        self._load()

        inn_q = _norm(inn)
        dose_mg = _parse_dose_mg(dosage)

        candidates: List[GrlsRecord] = []
        for rec in self._records:
            if inn_q and inn_q not in _norm(rec.inn):
                continue
            if not _dose_match(rec.forms, dose_mg):
                continue
            if not _form_match(rec.forms, dosage_form):
                continue
            candidates.append(rec)

        if not candidates:
            for rec in self._records:
                if inn_q and inn_q not in _norm(rec.inn):
                    continue
                if not _dose_match(rec.forms, dose_mg):
                    continue
                candidates.append(rec)

        best_by_trade: Dict[str, GrlsRecord] = {}
        for rec in candidates:
            key = rec.trade_name
            if key not in best_by_trade:
                best_by_trade[key] = rec
            else:
                if _date_sort_key(rec.reg_date_iso) < _date_sort_key(best_by_trade[key].reg_date_iso):
                    best_by_trade[key] = rec

        options = sorted(
            best_by_trade.values(),
            key=lambda x: (-self._score(x, inn_q, dose_mg, dosage_form), _date_sort_key(x.reg_date_iso), _norm(x.trade_name)),
        )

        options = options[:limit]

        default_trade = options[0].trade_name if options else ""
        return {
            "inn": inn,
            "dosage": dosage,
            "dosage_form": dosage_form,
            "default_trade_name": default_trade,
            "options": [
                {
                    "trade_name": r.trade_name,
                    "reg_no": r.reg_no,
                    "reg_date": r.reg_date_iso,
                    "forms": r.forms,
                    "inn": r.inn,
                    "score": self._score(r, inn_q, dose_mg, dosage_form),
                }
                for r in options
            ],
        }

    def search(
        self,
        query: str,
        dosage: str = "",
        dosage_form: str = "",
        limit: int = 20,
    ) -> Dict[str, Any]:
        self._load()
        q = _norm(query)
        dose_mg = _parse_dose_mg(dosage)
        matches: List[GrlsRecord] = []
        for rec in self._records:
            if q and (q not in _norm(rec.trade_name) and q not in _norm(rec.inn)):
                continue
            matches.append(rec)
        matches = sorted(
            matches,
            key=lambda r: (-self._score(r, q, dose_mg, dosage_form), _date_sort_key(r.reg_date_iso), _norm(r.trade_name)),
        )
        matches = matches[:limit]
        return {
            "query": query,
            "dosage": dosage,
            "dosage_form": dosage_form,
            "items": [
                {
                    "trade_name": r.trade_name,
                    "reg_no": r.reg_no,
                    "reg_date": r.reg_date_iso,
                    "forms": r.forms,
                    "inn": r.inn,
                    "score": self._score(r, q, dose_mg, dosage_form),
                }
                for r in matches
            ],
        }
